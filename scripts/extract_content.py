#!/usr/bin/env python3
"""Take stock of what a source document holds, so that a deck can carry all of it.

    python3 extract_content.py <document> <out-dir> [--max-rows N]

Works on PDF, DOCX, PPTX, XLSX and Markdown. Writes into <out-dir>:

  outline.json   the document's parts in order: {id, parent, level, title, page, words, tables,
                 charts, pictures}. Give it to deck.json as meta.coverage, and name the parts each
                 slide carries in that slide's "covers"; the renderer then reports every part with
                 substance that no slide carries and meta.omit does not explain.
  outline.md     the same, readable: one line per part, indented by level.
  tables/T01.json, T01.csv
                 every table, as "head" + "rows" ready for the table layout, with where it is.
  charts/C01.json
                 every native chart in a PPTX, DOCX or XLSX, with its data, ready for the chart
                 layout ("condition" is left as a placeholder: say on the slide what the data is).
  index.md       one line per table and chart with its first row, to pick from.

Pictures are not extracted here: scripts/extract_images.py does that.
PDF needs PyMuPDF (pip install pymupdf); XLSX needs openpyxl. Neither changes the document.
"""
import csv
import json
import os
import re
import sys
import zipfile
import posixpath
import xml.etree.ElementTree as ET

NS = {
    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
    'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'rel': 'http://schemas.openxmlformats.org/package/2006/relationships',
    's': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
    'xdr': 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing',
}
R_ID = '{%s}id' % NS['r']


def clean(t):
    return re.sub(r'\s+', ' ', str(t if t is not None else '')).strip()


def words(t):
    return len(re.findall(r'\w+', t or ''))


class Doc:
    """Collects sections, tables and charts in document order."""

    def __init__(self, source, kind):
        self.source, self.kind = source, kind
        self.sections, self.tables, self.charts = [], [], []
        self.stack = []  # (level, id)
        self.counters = []
        self.cur = None

    def section(self, title, level=1, page=None, sid=None, count=True):
        level = max(1, min(6, level))
        while len(self.counters) < level:
            self.counters.append(0)
        self.counters = self.counters[:level]
        if count:
            self.counters[level - 1] += 1
        while self.stack and self.stack[-1][0] >= level:
            self.stack.pop()
        parent = self.stack[-1][1] if self.stack else None
        sid = sid or '.'.join(str(n or 1) for n in self.counters)
        sec = {'id': sid, 'parent': parent, 'level': level, 'title': clean(title)[:160], 'page': page,
               'words': 0, 'tables': 0, 'charts': 0, 'pictures': 0}
        self.sections.append(sec)
        self.stack.append((level, sid))
        self.cur = sec
        return sec

    def ensure(self, page=None):
        if not self.cur:
            self.section('(opening, before the first heading)', 1, page, sid='0', count=False)
        return self.cur

    def text(self, t, page=None):
        self.ensure(page)['words'] += words(t)

    def picture(self, n=1, page=None):
        self.ensure(page)['pictures'] += n

    def table(self, rows, where, page=None):
        rows = [[clean(c) for c in r] for r in rows if r is not None]
        rows = [r for r in rows if any(r)]
        if len(rows) < 2:
            return
        n = max(len(r) for r in rows)
        rows = [r + [''] * (n - len(r)) for r in rows]
        keep = [j for j in range(n) if any(r[j] for r in rows)]
        rows = [[r[j] for j in keep] for r in rows]
        if len(keep) < 2:
            return
        sec = self.ensure(page)
        sec['tables'] += 1
        tid = 'T%02d' % (len(self.tables) + 1)
        self.tables.append({'id': tid, 'where': where, 'section': sec['id'], 'head': rows[0], 'rows': rows[1:],
                            'n_rows': len(rows) - 1, 'n_cols': len(keep)})

    def chart(self, ch, where, page=None):
        if not ch:
            return
        sec = self.ensure(page)
        sec['charts'] += 1
        cid = 'C%02d' % (len(self.charts) + 1)
        self.charts.append({'id': cid, 'where': where, 'section': sec['id'], 'chart': ch})


# ---------------------------------------------------------------- charts (OOXML)
CHART_TYPES = {'barChart': 'bar', 'bar3DChart': 'bar', 'lineChart': 'line', 'line3DChart': 'line', 'pieChart': 'pie',
               'pie3DChart': 'pie', 'ofPieChart': 'pie', 'doughnutChart': 'doughnut', 'areaChart': 'area', 'area3DChart': 'area',
               'scatterChart': 'scatter', 'radarChart': 'radar', 'bubbleChart': 'bubble'}


def _pts(node):
    if node is None:
        return []
    out = {}
    for cache in node.iter():
        tag = cache.tag.split('}')[-1]
        if tag in ('strCache', 'numCache', 'strLit', 'numLit'):
            for pt in cache.findall('c:pt', NS):
                v = pt.find('c:v', NS)
                out[int(pt.get('idx', len(out)))] = v.text if v is not None else ''
            if out:
                break
    if not out:
        return []
    return [out.get(i, '') for i in range(max(out) + 1)]


def parse_chart(xml):
    try:
        root = ET.fromstring(xml)
    except ET.ParseError:
        return None
    plot = root.find('.//c:plotArea', NS)
    if plot is None:
        return None
    title = ' '.join(t.text or '' for t in root.findall('./c:chart/c:title//a:t', NS)).strip()
    groups = [g for g in plot if g.tag.split('}')[-1] in CHART_TYPES]
    if not groups:
        return None
    g = groups[0]
    kind = CHART_TYPES[g.tag.split('}')[-1]]
    bar_dir = g.find('c:barDir', NS)
    grouping = g.find('c:grouping', NS)
    series, cats = [], []
    for grp in groups:
        for ser in grp.findall('c:ser', NS):
            name = ' '.join(v.text or '' for v in ser.findall('./c:tx//c:v', NS)).strip() or \
                ' '.join(t.text or '' for t in ser.findall('./c:tx//a:t', NS)).strip()
            c = _pts(ser.find('c:cat', NS)) or _pts(ser.find('c:xVal', NS))
            v = _pts(ser.find('c:val', NS)) or _pts(ser.find('c:yVal', NS))
            if not cats and c:
                cats = c
            vals = []
            for x in v:
                try:
                    vals.append(round(float(x), 6))
                except (TypeError, ValueError):
                    vals.append(0)
            series.append({'name': name or 'Series %d' % (len(series) + 1), 'values': vals})
    if not series:
        return None
    n = max(len(cats), max(len(s['values']) for s in series))
    cats = [clean(x) for x in cats] + [str(i + 1) for i in range(len(cats), n)]
    for s in series:
        s['values'] = s['values'] + [0] * (n - len(s['values']))
    ch = {'type': kind if kind in ('bar', 'line', 'pie', 'doughnut', 'area') else kind, 'title': title, 'categories': cats,
          'series': series, 'condition': '[say which data, setting and period this chart shows]'}
    if kind == 'bar' and bar_dir is not None and bar_dir.get('val') == 'bar':
        ch['horizontal'] = True
    if grouping is not None and grouping.get('val') in ('stacked', 'percentStacked'):
        ch['stacked'] = True
    if kind in ('scatter', 'radar', 'bubble'):
        ch['note'] = 'the chart layout has no %s type; show the data another way' % kind
    return ch


def rels(z, part):
    d, b = posixpath.split(part)
    rp = posixpath.join(d, '_rels', b + '.rels')
    out = {}
    if rp in z.namelist():
        for r in ET.fromstring(z.read(rp)).findall('rel:Relationship', NS):
            t = r.get('Target')
            out[r.get('Id')] = t if r.get('TargetMode') == 'External' else posixpath.normpath(posixpath.join(d, t))
    return out


# ---------------------------------------------------------------- DOCX
def from_docx(path, doc):
    z = zipfile.ZipFile(path)
    names = {}
    if 'word/styles.xml' in z.namelist():
        for st in ET.fromstring(z.read('word/styles.xml')).findall('w:style', NS):
            nm = st.find('w:name', NS)
            ol = st.find('.//w:outlineLvl', NS)
            names[st.get('{%s}styleId' % NS['w'])] = (nm.get('{%s}val' % NS['w']) if nm is not None else '',
                                                     int(ol.get('{%s}val' % NS['w'])) if ol is not None else None)
    rel = rels(z, 'word/document.xml')
    body = ET.fromstring(z.read('word/document.xml')).find('w:body', NS)
    V = '{%s}val' % NS['w']

    def ptext(p):
        return ''.join(t.text or '' for t in p.iter('{%s}t' % NS['w']))

    for el in body:
        tag = el.tag.split('}')[-1]
        if tag == 'p':
            t = clean(ptext(el))
            sty = el.find('w:pPr/w:pStyle', NS)
            ol = el.find('w:pPr/w:outlineLvl', NS)
            level = None
            if ol is not None and ol.get(V) and ol.get(V).isdigit() and int(ol.get(V)) < 9:
                level = int(ol.get(V)) + 1
            elif sty is not None:
                nm, lvl = names.get(sty.get(V), ('', None))
                m = re.match(r'(?i)^(heading|judul)\s*(\d)', nm or sty.get(V) or '')
                if m:
                    level = int(m.group(2))
                elif re.match(r'(?i)^title$', nm or ''):
                    level = 1
                elif lvl is not None and lvl < 9:
                    level = lvl + 1
            if level and t:
                doc.section(t, level)
            else:
                doc.text(t)
            pics = len(el.findall('.//a:blip', NS)) + len(el.findall('.//{urn:schemas-microsoft-com:vml}imagedata'))
            if pics:
                doc.picture(pics)
            for cref in el.findall('.//c:chart', NS):
                tgt = rel.get(cref.get(R_ID))
                if tgt and tgt in z.namelist():
                    doc.chart(parse_chart(z.read(tgt)), 'chart in "%s"' % (doc.cur['title'] if doc.cur else ''))
        elif tag == 'tbl':
            rows = []
            for tr in el.findall('w:tr', NS):
                row = []
                for tc in tr.findall('w:tc', NS):
                    row.append(' '.join(clean(ptext(p)) for p in tc.findall('.//w:p', NS)).strip())
                    span = tc.find('w:tcPr/w:gridSpan', NS)
                    if span is not None and span.get(V, '1').isdigit():
                        row.extend([''] * (int(span.get(V)) - 1))
                rows.append(row)
            doc.text(' '.join(' '.join(r) for r in rows))
            doc.table(rows, 'table in "%s"' % (doc.cur['title'] if doc.cur else 'opening'))


# ---------------------------------------------------------------- PPTX
def from_pptx(path, doc):
    z = zipfile.ZipFile(path)
    pres = ET.fromstring(z.read('ppt/presentation.xml'))
    prel = rels(z, 'ppt/presentation.xml')
    order = [prel[s.get(R_ID)] for s in pres.findall('p:sldIdLst/p:sldId', NS) if s.get(R_ID) in prel]
    for n, part in enumerate(order, 1):
        root = ET.fromstring(z.read(part))
        srel = rels(z, part)
        title = ''
        for sp in root.iter('{%s}sp' % NS['p']):
            ph = sp.find('.//p:nvPr/p:ph', NS)
            if ph is not None and ph.get('type') in ('title', 'ctrTitle'):
                title = clean(' '.join(t.text or '' for t in sp.iter('{%s}t' % NS['a'])))
                break
        texts = [clean(t.text) for t in root.iter('{%s}t' % NS['a'])]
        if not title:
            title = next((t for t in texts if len(t) > 3), '(slide without text)')
        doc.section(title, 1, page=n, sid='S%d' % n)
        doc.text(' '.join(texts), n)
        doc.picture(len(root.findall('.//p:pic', NS)), n)
        for gf in root.iter('{%s}graphicFrame' % NS['p']):
            tbl = gf.find('.//a:tbl', NS)
            if tbl is not None:
                rows = [[clean(' '.join(t.text or '' for t in tc.iter('{%s}t' % NS['a']))) for tc in tr.findall('a:tc', NS)]
                        for tr in tbl.findall('a:tr', NS)]
                doc.table(rows, 'slide %d' % n, n)
            cref = gf.find('.//c:chart', NS)
            if cref is not None:
                tgt = srel.get(cref.get(R_ID))
                if tgt and tgt in z.namelist():
                    doc.chart(parse_chart(z.read(tgt)), 'slide %d' % n, n)


# ---------------------------------------------------------------- XLSX
def cell(v):
    import datetime
    if v is None:
        return ''
    if isinstance(v, datetime.datetime):
        return v.date().isoformat() if (v.hour, v.minute, v.second) == (0, 0, 0) else v.strftime('%Y-%m-%d %H:%M')
    if isinstance(v, float):
        return '%g' % v if abs(v) < 1e15 else str(v)
    return str(v)


def from_xlsx(path, doc, max_rows):
    try:
        import openpyxl
    except ImportError:
        sys.exit('XLSX needs openpyxl: pip install openpyxl')
    z = zipfile.ZipFile(path)
    wb_rel = rels(z, 'xl/workbook.xml')
    wbx = ET.fromstring(z.read('xl/workbook.xml'))
    sheet_part = {s.get('name'): wb_rel.get(s.get(R_ID)) for s in wbx.findall('s:sheets/s:sheet', NS)}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for i, ws in enumerate(wb.worksheets, 1):
        doc.section(ws.title, 1, page=i, sid='S%d' % i)
        grid = []
        for k, row in enumerate(ws.iter_rows(values_only=True)):
            if k >= 5000:
                break
            grid.append([cell(v) for v in row])
        doc.cur['words'] = sum(1 for r in grid for v in r if v.strip())
        # blocks of rows separated by empty rows are separate tables
        block = []
        for r in grid + [[]]:
            if any(v.strip() for v in r):
                block.append(r)
            elif block:
                if len(block) >= 2:
                    doc.table(block[:max_rows + 1], 'sheet "%s"' % ws.title, i)
                    if len(block) > max_rows + 1:
                        doc.tables[-1]['truncated_from'] = len(block) - 1
                block = []
        part = sheet_part.get(ws.title)
        if part:
            for tgt in rels(z, part).values():
                if 'drawings/' in tgt and tgt in z.namelist():
                    for ct in rels(z, tgt).values():
                        if 'charts/' in ct and ct in z.namelist():
                            doc.chart(parse_chart(z.read(ct)), 'sheet "%s"' % ws.title, i)
    wb.close()


# ---------------------------------------------------------------- PDF
NUMBERED = re.compile(r'^(BAB\s+[IVXLC\d]+\b|[IVXLC]{1,5}\.\s|\d{1,2}(\.\d{1,2}){0,3}\.?\s|[A-Z]\.\s)')


def from_pdf(path, doc):
    try:
        import fitz
    except ImportError:
        sys.exit('PDF needs PyMuPDF: pip install pymupdf')
    pdf = fitz.open(path)
    lines = []  # (page, y, text, size, bold)
    for pno, page in enumerate(pdf, 1):
        for b in page.get_text('dict')['blocks']:
            for ln in b.get('lines', []):
                spans = [s for s in ln['spans'] if s['text'].strip()]
                if not spans:
                    continue
                t = clean(''.join(s['text'] for s in spans))
                size = max(s['size'] for s in spans)
                bold = all((s['flags'] & 16) or 'Bold' in s['font'] for s in spans)
                lines.append((pno, ln['bbox'][1], t, round(size, 1), bold))
    # repeated running heads and feet are not content
    freq = {}
    for l in lines:
        freq[l[2]] = freq.get(l[2], set()) | {l[0]}
    repeat = {t for t, ps in freq.items() if len(ps) > max(2, len(pdf) * 0.4)}
    sizes = {}
    for l in lines:
        sizes[l[3]] = sizes.get(l[3], 0) + len(l[2])
    body = max(sizes, key=sizes.get) if sizes else 10
    heads = {}
    toc = pdf.get_toc(simple=True)
    if len(toc) >= 2:
        for lvl, title, page in toc:
            key = clean(title).lower()[:40]
            for i, l in enumerate(lines):
                if l[0] == page and i not in heads and l[2].lower().startswith(key[:25]):
                    heads[i] = (lvl, title)
                    break
            else:
                cand = [i for i, l in enumerate(lines) if l[0] == page]
                if cand:
                    heads.setdefault(cand[0], (lvl, title))
    else:
        big = sorted({l[3] for l in lines if l[3] >= body + 1.5}, reverse=True)
        # a table of contents is not a list of headings
        toc_pages = {l[0] for l in lines if re.match(r'(?i)^(daftar isi|table of contents|contents|isi)$', l[2])}
        toc_pages |= {pg for pg in {l[0] for l in lines} if sum(1 for l in lines if l[0] == pg and re.search(r'(\.{4,}|…{2,})\s*\d+$', l[2])) >= 4}
        prev = None
        for i, l in enumerate(lines):
            t = l[2]
            if t in repeat or l[0] in toc_pages or len(t) > 110 or len(t) < 2 or re.match(r'^[\d\s.,/-]+$', t):
                prev = None
                continue
            m = NUMBERED.match(t + ' ')
            if m and len(t[m.end() - 1:].strip()) < 2:
                prev = None
                continue
            lvl = None
            if l[3] >= body + 1.5:
                lvl = min(3, big.index(l[3]) + 1) if l[3] in big else 1
            elif l[4] and m and len(t) <= 90:
                depth = t.split()[0].rstrip('.').count('.') + 1 if re.match(r'\d', t) else 1
                lvl = min(3, depth + (1 if big else 0))
            if lvl is None:
                prev = None
                continue
            # a heading that runs over two or more lines is one heading
            if prev is not None and lines[prev][0] == l[0] and lines[prev][3] == l[3] and not m and (prev == i - 1 or 0 < l[1] - lines[prev][1] < l[3] * 2.2):
                heads[prev] = (heads[prev][0], heads[prev][1] + ' ' + t)
                continue
            heads[i] = (lvl, t)
            prev = i
    # tables and pictures, placed by page and height
    marks = []
    for pno, page in enumerate(pdf, 1):
        try:
            for tab in page.find_tables().tables:
                marks.append((pno, tab.bbox[1], 'table', tab.extract()))
        except Exception:
            pass
        for img in page.get_images(full=True):
            try:
                rects = page.get_image_rects(img[0])
            except Exception:
                rects = []
            r = rects[0] if rects else None
            if r is not None and r.width * r.height < 2500:  # icons and bullets
                continue
            marks.append((pno, r.y0 if r is not None else 0, 'picture', None))
    marks.sort(key=lambda m: (m[0], m[1]))
    mi = 0
    for i, l in enumerate(lines):
        while mi < len(marks) and (marks[mi][0], marks[mi][1]) < (l[0], l[1]):
            m = marks[mi]
            if m[2] == 'table':
                doc.table(m[3], 'page %d' % m[0], m[0])
            else:
                doc.picture(1, m[0])
            mi += 1
        if i in heads:
            lvl, title = heads[i]
            doc.section(title, lvl, page=l[0])
        elif l[2] not in repeat and not any(l[2] in h[1] and abs(j - i) < 4 for j, h in heads.items()):
            doc.text(l[2], l[0])
    for m in marks[mi:]:
        if m[2] == 'table':
            doc.table(m[3], 'page %d' % m[0], m[0])
        else:
            doc.picture(1, m[0])


# ---------------------------------------------------------------- Markdown
def from_md(path, doc):
    lines = open(path, encoding='utf-8').read().splitlines()
    i, fence = 0, False
    while i < len(lines):
        ln = lines[i]
        if ln.strip().startswith('```'):
            fence = not fence
            i += 1
            continue
        if fence:
            doc.text(ln)
            i += 1
            continue
        m = re.match(r'^(#{1,6})\s+(.*)$', ln)
        if m:
            doc.section(re.sub(r'[*_`]', '', m.group(2)), len(m.group(1)))
            i += 1
            continue
        if ln.strip().startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.match(r'^:?-{2,}:?$', c) for c in cells if c):
                    rows.append([re.sub(r'(\*\*|__|`)', '', c) for c in cells])
                i += 1
            doc.text(' '.join(' '.join(r) for r in rows))
            doc.table(rows, 'table in "%s"' % (doc.cur['title'] if doc.cur else 'opening'))
            continue
        pics = len(re.findall(r'!\[[^\]]*\]\([^)]+\)', ln))
        if pics:
            doc.picture(pics)
        doc.text(re.sub(r'!\[[^\]]*\]\([^)]+\)', '', ln))
        i += 1


# ---------------------------------------------------------------- output
def write(doc, out):
    os.makedirs(os.path.join(out, 'tables'), exist_ok=True)
    os.makedirs(os.path.join(out, 'charts'), exist_ok=True)
    with open(os.path.join(out, 'outline.json'), 'w', encoding='utf-8') as f:
        json.dump({'source': doc.source, 'kind': doc.kind, 'sections': doc.sections}, f, ensure_ascii=False, indent=1)
    with open(os.path.join(out, 'outline.md'), 'w', encoding='utf-8') as f:
        f.write('# %s\n\nid · title · where · words · tables · charts · pictures\n\n' % doc.source)
        for s in doc.sections:
            bits = ['p. %s' % s['page'] if s['page'] else '', '%d words' % s['words']]
            bits += ['%d %s' % (s[k], k) for k in ('tables', 'charts', 'pictures') if s[k]]
            f.write('%s- **%s** %s — %s\n' % ('  ' * (s['level'] - 1), s['id'], s['title'], ' · '.join(b for b in bits if b)))
    for t in doc.tables:
        with open(os.path.join(out, 'tables', t['id'] + '.json'), 'w', encoding='utf-8') as f:
            json.dump(t, f, ensure_ascii=False, indent=1)
        with open(os.path.join(out, 'tables', t['id'] + '.csv'), 'w', encoding='utf-8', newline='') as f:
            csv.writer(f).writerows([t['head']] + t['rows'])
    for c in doc.charts:
        with open(os.path.join(out, 'charts', c['id'] + '.json'), 'w', encoding='utf-8') as f:
            json.dump(c, f, ensure_ascii=False, indent=1)
    with open(os.path.join(out, 'index.md'), 'w', encoding='utf-8') as f:
        f.write('# Tables and charts in %s\n\n' % doc.source)
        for t in doc.tables:
            f.write('- **%s** (%s, part %s) %d rows × %d columns%s — %s\n' % (
                t['id'], t['where'], t['section'], t['n_rows'], t['n_cols'],
                ', cut from %d rows' % t['truncated_from'] if t.get('truncated_from') else '', ' | '.join(t['head'])[:160]))
        for c in doc.charts:
            ch = c['chart']
            f.write('- **%s** (%s, part %s) %s chart "%s": %d categories × %d series%s\n' % (
                c['id'], c['where'], c['section'], ch['type'], ch.get('title', ''), len(ch['categories']), len(ch['series']),
                ' — ' + ch['note'] if ch.get('note') else ''))


def main():
    import argparse
    ap = argparse.ArgumentParser(description='Outline, tables and charts of a PDF, DOCX, PPTX, XLSX or Markdown file.')
    ap.add_argument('document')
    ap.add_argument('out_dir')
    ap.add_argument('--max-rows', type=int, default=200, help='rows kept per spreadsheet table (default 200)')
    a = ap.parse_args()
    src, out = a.document, a.out_dir
    ext = os.path.splitext(src)[1].lower()
    doc = Doc(os.path.basename(src), ext.lstrip('.'))
    if ext == '.pdf':
        from_pdf(src, doc)
    elif ext == '.docx':
        from_docx(src, doc)
    elif ext == '.pptx':
        from_pptx(src, doc)
    elif ext in ('.xlsx', '.xlsm'):
        from_xlsx(src, doc, a.max_rows)
    elif ext in ('.md', '.markdown'):
        from_md(src, doc)
    else:
        sys.exit('unsupported file type: %s (PDF, DOCX, PPTX, XLSX, MD)' % ext)
    write(doc, out)
    subst = sum(1 for s in doc.sections if s['words'] >= 40 or s['tables'] or s['charts'] or s['pictures'])
    print('%s: %d parts (%d with substance), %d tables, %d charts, %d pictures -> %s' % (
        doc.source, len(doc.sections), subst, len(doc.tables), len(doc.charts), sum(s['pictures'] for s in doc.sections), out))
    print('read %s/outline.md and %s/index.md' % (out, out))


if __name__ == '__main__':
    main()
